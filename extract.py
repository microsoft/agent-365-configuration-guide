import os, json, io
from pptx import Presentation
from pptx.util import Emu

OUT = os.path.dirname(os.path.abspath(__file__))
# Path to the source PowerPoint deck (not included in this repository).
SRC = os.environ.get("DECK_PPTX", os.path.join(OUT, "Agent365 Config Guide.pptx"))
IMG = os.path.join(OUT, "img")
os.makedirs(IMG, exist_ok=True)

prs = Presentation(SRC)
data = []

def walk(shapes, acc, prefix=""):
    for sh in shapes:
        item = {"name": sh.shape_name if hasattr(sh, 'shape_name') else sh.name,
                "type": str(sh.shape_type),
                "left": sh.left, "top": sh.top, "w": sh.width, "h": sh.height}
        if sh.shape_type == 6:  # group
            item["kind"] = "group"
            acc.append(item)
            walk(sh.shapes, acc, prefix + "  ")
            continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            item["kind"] = "text"
            item["text"] = sh.text_frame.text
        elif sh.shape_type == 13 or sh.__class__.__name__ == "Picture":
            item["kind"] = "picture"
        elif getattr(sh, "has_table", False) and sh.has_table:
            item["kind"] = "table"
            item["table"] = [[c.text for c in r.cells] for r in sh.table.rows]
        else:
            item["kind"] = "other"
        acc.append(item)

for i, slide in enumerate(prs.slides, 1):
    acc = []
    walk(slide.shapes, acc)
    notes = ""
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
        notes = slide.notes_slide.notes_text_frame.text
    data.append({"slide": i, "shapes": acc, "notes": notes})

with open(os.path.join(OUT, "deck.json"), "w", encoding="utf-8") as f:
    json.dump(data, f, ensure_ascii=False, indent=1)

print("slides:", len(prs.slides), "size:", prs.slide_width, prs.slide_height)
